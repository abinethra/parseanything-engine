import csv
import os
from typing import Any, Dict, List
import docx
import openpyxl
from pptx import Presentation

def parse_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extracts structured text blocks from Word documents."""
    blocks = []
    doc = docx.Document(file_path)
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        
        block_type = "heading" if para.style.name.startswith("Heading") else "paragraph"
        blocks.append({
            "block_id": f"docx_p{idx + 1}",
            "page_number": 1,
            "text": text,
            "block_type": block_type,
            "bbox": {"x0": 0.0, "y0": 0.0, "x1": 0.0, "y1": 0.0},
            "confidence": 1.0,
            "flags": ["docx_source"]
        })
    return blocks

def parse_pptx(file_path: str) -> List[Dict[str, Any]]:
    """Extracts slide text blocks from PowerPoint presentations."""
    blocks = []
    prs = Presentation(file_path)
    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            if hasattr(shape, "text") and shape.text.strip():
                blocks.append({
                    "block_id": f"s{slide_idx + 1}_sh{shape_idx + 1}",
                    "page_number": slide_idx + 1,
                    "text": shape.text.strip(),
                    "block_type": "paragraph",
                    "bbox": {"x0": 0.0, "y0": 0.0, "x1": 0.0, "y1": 0.0},
                    "confidence": 1.0,
                    "flags": ["pptx_source"]
                })
    return blocks

def parse_xlsx(file_path: str) -> List[Dict[str, Any]]:
    """Extracts tables from Excel spreadsheets into Markdown tables."""
    blocks = []
    wb = openpyxl.load_workbook(file_path, data_only=True)
    for sheet_idx, sheet_name in enumerate(wb.sheetnames):
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
            
        markdown_lines = []
        headers = [str(cell or "").strip() for cell in rows[0]]
        markdown_lines.append("| " + " | ".join(headers) + " |")
        markdown_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
        
        for row in rows[1:]:
            cells = [str(cell or "").strip().replace("\n", " ") for cell in row]
            markdown_lines.append("| " + " | ".join(cells) + " |")
            
        blocks.append({
            "block_id": f"sheet_{sheet_idx + 1}",
            "page_number": sheet_idx + 1,
            "text": "\n".join(markdown_lines),
            "block_type": "table",
            "bbox": {"x0": 0.0, "y0": 0.0, "x1": 0.0, "y1": 0.0},
            "confidence": 1.0,
            "flags": [f"sheet_name_{sheet_name}"]
        })
    return blocks
